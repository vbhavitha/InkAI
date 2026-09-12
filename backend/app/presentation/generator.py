from pathlib import Path

from pptx import Presentation

from .slides import (
    add_slide_title,
    add_slide_content,
)


def generate_presentation(
    title: str,
    slides: list[dict],
    output_path: str | Path,
):
    output_path = Path(output_path)

    output_path.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    presentation = Presentation()

    # ----------------------------------------------------------
    # TITLE SLIDE
    # ----------------------------------------------------------

    title_layout = presentation.slide_layouts[0]

    title_slide = presentation.slides.add_slide(
        title_layout
    )

    title_slide.shapes.title.text = title

    # ----------------------------------------------------------
    # CONTENT SLIDES
    # ----------------------------------------------------------

    blank_layout = presentation.slide_layouts[6]

    for slide_data in slides:
        slide = presentation.slides.add_slide(
            blank_layout
        )

        add_slide_title(
            slide,
            slide_data.get(
                "title",
                "Untitled",
            ),
        )

        add_slide_content(
            slide,
            slide_data.get(
                "content",
                [],
            ),
        )

    presentation.save(
        str(output_path)
    )

    return output_path