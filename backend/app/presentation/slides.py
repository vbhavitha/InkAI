from pptx.util import Pt

from .layouts import (
    TITLE_LEFT,
    TITLE_TOP,
    TITLE_WIDTH,
    TITLE_HEIGHT,
    CONTENT_LEFT,
    CONTENT_TOP,
    CONTENT_WIDTH,
    CONTENT_HEIGHT,
)
from .themes import DEFAULT_THEME


def add_slide_title(
    slide,
    title: str,
):
    textbox = slide.shapes.add_textbox(
        TITLE_LEFT,
        TITLE_TOP,
        TITLE_WIDTH,
        TITLE_HEIGHT,
    )

    frame = textbox.text_frame

    paragraph = frame.paragraphs[0]

    paragraph.text = title

    paragraph.font.name = DEFAULT_THEME["font"]

    paragraph.font.size = Pt(28)

    paragraph.font.bold = True

    return textbox


def add_slide_content(
    slide,
    content: list[str],
):
    textbox = slide.shapes.add_textbox(
        CONTENT_LEFT,
        CONTENT_TOP,
        CONTENT_WIDTH,
        CONTENT_HEIGHT,
    )

    frame = textbox.text_frame

    frame.clear()

    for index, item in enumerate(content):
        paragraph = (
            frame.paragraphs[0]
            if index == 0
            else frame.add_paragraph()
        )

        paragraph.text = item

        paragraph.font.name = (
            DEFAULT_THEME["font"]
        )

        paragraph.font.size = Pt(20)

        paragraph.space_after = Pt(12)

        paragraph.level = 0

    return textbox